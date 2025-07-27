from sshtunnel import SSHTunnelForwarder
import os
import json
from pyspark.sql import SparkSession
from openpyxl import load_workbook
from concurrent.futures import ThreadPoolExecutor, as_completed
import logging
import datetime
import decimal

def json_serial(obj):
    """JSON serializer for objects not serializable by default json code"""
    if isinstance(obj, datetime.datetime) or isinstance(obj, datetime.date):
        return obj.isoformat()
    if isinstance(obj, decimal.Decimal):
        return float(obj)
    return str(obj)

logging.basicConfig(filename="/home/vagrant/datalake-mavis/provision/logs/datalake.log", level=logging.INFO)

class DateTimeEncoder(json.JSONEncoder):
    def default(self, obj):
        if isinstance(obj, datetime):
            return obj.isoformat()  # Convertir datetime en chaîne ISO
        return super().default(obj)
    
def discover_postgres(source, spark, source_index):
    logging.info(f"Découverte pour {source['name']}")
    ssh_cfg = source.get("ssh")
    db_cfg = source["db"]
    tables_info = []

    try:
        ssh_password = os.getenv(f"SSH_PASSWORD_{source['name']}", ssh_cfg.get("password", "0000!"))
        db_password = os.getenv(f"POSTGRES_PASSWORD_{source['name']}", db_cfg.get("password", "00000!"))
        logging.info(f"SSH Password defined: {bool(ssh_password)}")
        logging.info(f"DB Password defined: {bool(db_password)}")

        with SSHTunnelForwarder(
            (ssh_cfg["host"], ssh_cfg["port"]),
            ssh_username=ssh_cfg["user"],
            ssh_password=ssh_password,
            remote_bind_address=(db_cfg["host"], db_cfg["port"]),
            local_bind_address=('127.0.0.1', 15432 + source_index)
        ) as tunnel:
            logging.info(f"Tunnel établi sur port local {tunnel.local_bind_port}")
            url = f"jdbc:postgresql://127.0.0.1:{tunnel.local_bind_port}/{db_cfg['database']}"
            properties = {
                "user": db_cfg["user"],
                "password": db_password,
                "driver": "org.postgresql.Driver",
                "ssl": "false"
            }

            # Récupérer les tables
            logging.info("Exécution de la requête pour récupérer les tables...")
            tables_to_include = source.get("tables_to_include", [])

            if tables_to_include:
                # Format SQL : 'table1','table2',...
                tables_str = ",".join(f"'{t}'" for t in tables_to_include)
                query = f"""
                    (SELECT table_name, table_type 
                    FROM information_schema.tables 
                    WHERE table_schema = 'public' 
                    AND table_type IN ('BASE TABLE', 'VIEW')
                    AND table_name IN ({tables_str})
                    ) t
                """
            else:
                query = """
                    (SELECT table_name, table_type 
                    FROM information_schema.tables 
                    WHERE table_schema = 'public' 
                    AND table_type IN ('BASE TABLE', 'VIEW')
                    ) t
                """
            df_tables = spark.read.jdbc(
                url=url,
                table=query,
                properties=properties
            )
            tables = [{"table_name": row["table_name"], "table_type": row["table_type"]} for row in df_tables.collect()]
            logging.info(f"Nombre de tables trouvées : {len(tables)}")
            logging.info(f"Tables : {tables}")

            for table in tables:
                table_name = table["table_name"]
                logging.info(f"Traitement de la table : {table_name}")
                try:
                    df_cols = spark.read.jdbc(
                        url=url,
                        table=f"(SELECT column_name, data_type FROM information_schema.columns WHERE table_name = '{table_name}' ORDER BY ordinal_position) t",
                        properties=properties
                    )
                    df_type = spark.read.jdbc(
                        url=url,
                        table=f"(SELECT table_type FROM information_schema.tables WHERE table_name = '{table_name}') t",
                        properties=properties
                    )
                    table_type = df_type.collect()[0]["table_type"] if df_type.count() > 0 else "UNKNOWN"
                    columns = [{"name": row["column_name"], "type": str(row["data_type"])} for row in df_cols.collect()]
                    logging.info(f"Colonnes pour {table_name} : {columns}")
                    df_data = spark.read.jdbc(url=url, table=table_name, properties=properties)
                    sample_data = df_data.limit(10).collect()
                    df_data.write.mode("overwrite").parquet(f"hdfs://localhost:9000/datalake/raw/{source['name']}/{table_name}")
                    tables_info.append({
                        "table_name": table_name,
                        "table_type": table_type,
                        "columns": columns,
                        "row_count": df_data.count(),
                        "sample_data": [row.asDict() for row in sample_data]
                    })
                    logging.info(f"Table {table_name} processed: {len(columns)} columns, {df_data.count()} rows")
                except Exception as e:
                    logging.error(f"Error processing table {table_name}: {str(e)}")
                    continue
    except Exception as e:
        logging.error(f"Error connecting to {source['name']}: {str(e)}")
        return tables_info

    logging.info(f"Tables info collectées : {tables_info}")
    return tables_info

def discover_excel(source, spark):
    df = spark.read.format("com.crealytics.spark.excel") \
        .option("header", source["options"]["header"]) \
        .option("inferSchema", source["options"]["inferSchema"]) \
        .option("dataAddress", "PatientList") \
        .load(source["path"])
    columns = [{"name": col, "type": str(df.schema[col].dataType)} for col in df.columns]
    return [{"table_name": "PatientList", "columns": columns}]

def main():
    # Configuration Spark avec vos JARs explicites
    spark = SparkSession.builder \
        .appName("PivotSchemaDiscovery") \
        .master("local[*]") \
        .config("spark.jars", ",".join([
            "/home/vagrant/spark/jars/postgresql-42.7.3.jar",
            "/home/vagrant/spark/jars/spark-excel_2.12-3.5.0_0.20.3.jar",
            "/home/vagrant/spark/jars/poi-5.2.3.jar",
            "/home/vagrant/spark/jars/poi-ooxml-5.2.3.jar",
            "/home/vagrant/spark/jars/poi-ooxml-lite-5.2.3.jar",
            "/home/vagrant/spark/jars/xmlbeans-5.1.1.jar",
            "/home/vagrant/spark/jars/commons-compress-1.26.2.jar",
            "/home/vagrant/spark/jars/commons-collections4-4.4.jar",
            "/home/vagrant/spark/jars/ooxml-schemas-1.4.jar"
        ])) \
        .config("spark.executor.memory", "8g") \
        .config("spark.executor.cores", "4") \
        .config("spark.driver.memory", "4g") \
        .config("spark.sql.debug.maxToStringFields", "1000") \
        .getOrCreate()
        
    spark.sparkContext.setLogLevel("ERROR")

    # Charger data_sources.json depuis le dossier synchronisé
    with open("/home/vagrant/datalake-mavis/provision/config/data_sources.json") as f:
        sources = json.load(f)

    discovery_report = {}

    # Paralléliser les bases PostgreSQL avec ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=10) as executor:
        future_to_source = {
            executor.submit(discover_postgres, source, spark, i): source
            for i, source in enumerate(sources) if source["type"] == "postgres"
        }
        for future in as_completed(future_to_source):
            source = future_to_source[future]
            try:
                discovery_report[source["name"]] = future.result()
                print(f"✅ Découverte réussie pour {source['name']} (postgres)")
            except Exception as e:
                discovery_report[source["name"]] = {"error": str(e)}
                print(f"❌ Erreur pour {source['name']} (postgres): {str(e)}")
                logging.error(f"Failed for {source['name']}: {str(e)}")

    # Traiter les sources Excel séquentiellement
    for source in sources:
        if source["type"] == "excel":
            try:
                discovery_report[source["name"]] = discover_excel(source, spark)
                print(f"✅ Découverte réussie pour {source['name']} (excel)")
            except Exception as e:
                discovery_report[source["name"]] = {"error": str(e)}
                print(f"❌ Erreur pour {source['name']} (excel): {str(e)}")
                logging.error(f"Failed for {source['name']}: {str(e)}")

    # Créer le dossier de sortie et enregistrer le rapport
    output_dir = "/home/vagrant/datalake-mavis/provision/spark-jobs/discovery"
    os.makedirs(output_dir, exist_ok=True)
    output_file = os.path.join(output_dir, "discovery_report.json")
    with open(output_file, "w") as f:
        json.dump(discovery_report, f, indent=2, default=json_serial)

    print(f"\n✅ Rapport de découverte généré : {output_file}")
    spark.stop()

if __name__ == "__main__":
    main()
    